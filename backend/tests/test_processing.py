import io
import os
import tempfile
from pathlib import Path
import pytest
from PIL import Image
import docx
from pypdf import PdfWriter
from app.services.document_processing.pipeline import document_pipeline
from app.services.document_processing.text_processor import TextProcessor
from app.services.document_processing.docx_processor import DocxProcessor
from app.services.document_processing.image_processor import ImageProcessor
from app.services.document_processing.pdf_processor import PDFProcessor
from tests.test_documents import create_test_user

def test_text_processor_plain_text():
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as f:
        f.write("# Quarterly Financial Report\n\nRevenue grew by 25% year over year.\n\nOperating expenses remained stable.")
        tmp_path = Path(f.name)

    try:
        processor = TextProcessor()
        doc = processor.process(tmp_path, document_id="doc-123", document_type="P_AND_L")
        assert doc.document_id == "doc-123"
        assert doc.page_count == 1
        assert doc.total_text_blocks == 3
        assert doc.pages[0].text_blocks[0].block_type == "heading"
        assert "Quarterly Financial Report" in doc.pages[0].text_blocks[0].text
        assert "Revenue grew" in doc.full_text
    finally:
        if tmp_path.exists():
            os.remove(tmp_path)

def test_text_processor_csv_tables():
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w", encoding="utf-8") as f:
        f.write("Player,Runs,Wickets\nRohit,120,0\nVirat,85,1\nBumrah,10,4\n")
        tmp_path = Path(f.name)

    try:
        processor = TextProcessor()
        doc = processor.process(tmp_path, document_id="doc-csv", document_type="CRICKET_BROCHURE")
        assert doc.total_tables == 1
        table = doc.pages[0].tables[0]
        assert table.headers == ["Player", "Runs", "Wickets"]
        assert len(table.rows) == 4
        assert table.rows[1].cells[0].content == "Rohit"
        assert table.rows[2].cells[1].content == "85"
    finally:
        if tmp_path.exists():
            os.remove(tmp_path)

def test_text_processor_json():
    with tempfile.NamedTemporaryFile(suffix=".json", delete=False, mode="w", encoding="utf-8") as f:
        f.write('[{"item": "Server", "cost": 1200}, {"item": "Database", "cost": 800}]')
        tmp_path = Path(f.name)

    try:
        processor = TextProcessor()
        doc = processor.process(tmp_path, document_id="doc-json", document_type="INVOICE")
        assert doc.total_tables == 1
        assert doc.pages[0].tables[0].headers == ["item", "cost"]
        assert len(doc.pages[0].tables[0].rows) == 2
        assert "Server" in doc.full_text
    finally:
        if tmp_path.exists():
            os.remove(tmp_path)

def test_docx_processor_structure():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        tmp_path = Path(f.name)

    doc_obj = docx.Document()
    doc_obj.add_heading("Annual Balance Sheet", level=1)
    doc_obj.add_paragraph("Assets exceeded liabilities by $5M.")
    table = doc_obj.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Asset"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Cash"
    table.rows[1].cells[1].text = "5000000"
    doc_obj.save(str(tmp_path))

    try:
        processor = DocxProcessor()
        doc = processor.process(tmp_path, document_id="doc-docx", document_type="BALANCE_SHEET")
        assert doc.total_text_blocks == 2
        assert doc.total_tables == 1
        assert doc.pages[0].tables[0].headers == ["Asset", "Value"]
        assert "Annual Balance Sheet" in doc.full_text
    finally:
        if tmp_path.exists():
            os.remove(tmp_path)

def test_image_processor_execution():
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        tmp_path = Path(f.name)

    img = Image.new("RGB", (200, 100), color=(255, 255, 255))
    img.save(str(tmp_path))

    try:
        processor = ImageProcessor()
        doc = processor.process(tmp_path, document_id="doc-img", document_type="GENERAL")
        assert doc.page_count == 1
        assert doc.pages[0].width == 200.0
        assert doc.pages[0].height == 100.0
        assert doc.processor == "pytesseract_ocr_vision"
    finally:
        if tmp_path.exists():
            os.remove(tmp_path)

def test_pdf_processor_with_generated_pdf():
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        tmp_path = Path(f.name)

    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    with open(str(tmp_path), "wb") as f_out:
        writer.write(f_out)

    try:
        processor = PDFProcessor()
        doc = processor.process(tmp_path, document_id="doc-pdf", document_type="GENERAL")
        assert doc.page_count == 1
        assert doc.pages[0].width == 600.0
        assert doc.pages[0].height == 800.0
        assert doc.processor == "pdfplumber_layout_parser"
    finally:
        if tmp_path.exists():
            os.remove(tmp_path)

def test_pipeline_routing():
    p_pdf = document_pipeline.get_processor_for_file(Path("doc.pdf"))
    p_img = document_pipeline.get_processor_for_file(Path("photo.png"))
    p_txt = document_pipeline.get_processor_for_file(Path("notes.txt"))
    p_docx = document_pipeline.get_processor_for_file(Path("summary.docx"))

    assert isinstance(p_pdf, PDFProcessor)
    assert isinstance(p_img, ImageProcessor)
    assert isinstance(p_txt, TextProcessor)
    assert isinstance(p_docx, DocxProcessor)

    with pytest.raises(ValueError):
        document_pipeline.get_processor_for_file(Path("archive.zip"))

def test_document_process_and_content_endpoints_flow(client):
    user, token = create_test_user("procuser@intellirag.ai")
    sample_text = b"# INVOICE 1001\n\nTotal Amount Due: $4,500.00 USD\n\nDue Date: 2026-10-15"
    files = {"file": ("invoice_1001.txt", io.BytesIO(sample_text), "text/plain")}
    data = {"document_type": "INVOICE"}

    upload_res = client.post(
        "/api/documents/upload",
        files=files,
        data=data,
        headers={"Authorization": f"Bearer {token}"}
    )
    assert upload_res.status_code == 201
    doc_id = upload_res.json()["id"]
    assert upload_res.json()["status"] == "UPLOADED"

    proc_res = client.post(
        f"/api/documents/{doc_id}/process",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert proc_res.status_code == 200
    proc_data = proc_res.json()
    assert proc_data["status"] == "PROCESSED"
    assert proc_data["document_id"] == doc_id
    assert "Total Amount Due" in proc_data["extracted_text"]
    assert proc_data["extracted_metadata"]["page_count"] == 1
    assert len(proc_data["extracted_metadata"]["pages"][0]["text_blocks"]) >= 2

    content_res = client.get(
        f"/api/documents/{doc_id}/content",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert content_res.status_code == 200
    content_data = content_res.json()
    assert content_data["status"] == "PROCESSED"
    assert content_data["extracted_text"] == proc_data["extracted_text"]

def test_cross_user_processing_isolation(client):
    user1, token1 = create_test_user("proc1@intellirag.ai")
    user2, token2 = create_test_user("proc2@intellirag.ai")

    upload_res = client.post(
        "/api/documents/upload",
        files={"file": ("secret.txt", io.BytesIO(b"Confidential Report"), "text/plain")},
        headers={"Authorization": f"Bearer {token1}"}
    )
    doc_id = upload_res.json()["id"]

    assert client.post(
        f"/api/documents/{doc_id}/process",
        headers={"Authorization": f"Bearer {token2}"}
    ).status_code == 404

    assert client.get(
        f"/api/documents/{doc_id}/content",
        headers={"Authorization": f"Bearer {token2}"}
    ).status_code == 404

def test_unauthenticated_processing_endpoints(client):
    fake_id = "00000000-0000-0000-0000-000000000000"
    assert client.post(f"/api/documents/{fake_id}/process").status_code == 401
    assert client.get(f"/api/documents/{fake_id}/content").status_code == 401
