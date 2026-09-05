import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import List
import docx
from app.schemas.processing import (
    ExtractedDocument,
    DocumentPage,
    TextBlock,
    TableBlock,
    TableRow,
    TableCell,
)
from app.services.document_processing.base import BaseDocumentProcessor

class DocxProcessor(BaseDocumentProcessor):
    def process(
        self,
        file_path: Path,
        document_id: str,
        document_type: str = "GENERAL"
    ) -> ExtractedDocument:
        doc = docx.Document(file_path)
        text_blocks: List[TextBlock] = []
        table_blocks: List[TableBlock] = []
        all_text_parts: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style and p.style.name else ""
            is_heading = "heading" in style_name or "title" in style_name or (len(text) < 80 and text.isupper())
            b_type = "heading" if is_heading else "paragraph"

            text_blocks.append(
                TextBlock(
                    block_id=str(uuid.uuid4()),
                    block_type=b_type,
                    text=text,
                    page_number=1,
                    confidence=100.0
                )
            )
            all_text_parts.append(text)

        for t_idx, table in enumerate(doc.tables, start=1):
            rows: List[TableRow] = []
            headers: List[str] = []
            for r_idx, row in enumerate(table.rows):
                row_cells = [c.text.strip() for c in row.cells]
                if r_idx == 0:
                    headers = row_cells
                cells = [
                    TableCell(row_index=r_idx, col_index=c_idx, content=val)
                    for c_idx, val in enumerate(row_cells)
                ]
                rows.append(TableRow(row_index=r_idx, cells=cells))

            if rows:
                table_blocks.append(
                    TableBlock(
                        table_index=t_idx,
                        page_number=1,
                        headers=headers,
                        rows=rows
                    )
                )

        full_text = "\n\n".join(all_text_parts)

        page = DocumentPage(
            page_number=1,
            width=612.0,
            height=792.0,
            text_blocks=text_blocks,
            tables=table_blocks,
            has_ocr=False
        )

        return ExtractedDocument(
            document_id=document_id,
            processor="python_docx_layout_parser",
            page_count=1,
            total_text_blocks=len(text_blocks),
            total_tables=len(table_blocks),
            full_text=full_text,
            pages=[page],
            metadata={
                "document_type": document_type,
                "paragraph_count": len(text_blocks),
                "table_count": len(table_blocks),
            },
            processed_at=datetime.now(timezone.utc)
        )
